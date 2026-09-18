import json
import os
import tempfile
import unittest

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from checker import check_document, get_run_font_info, load_rules_from_json


class CheckerProfileTests(unittest.TestCase):
    def _document_path(self, folder):
        path = os.path.join(folder, "sample.docx")
        doc = Document()
        doc.add_paragraph("这是一段用于验证规则隔离的普通文字。")
        doc.save(path)
        return path

    def test_empty_enabled_checks_disables_thesis_specific_checks(self):
        with tempfile.TemporaryDirectory() as folder:
            document_path = self._document_path(folder)
            rules_path = os.path.join(folder, "generic.json")
            with open(rules_path, "w", encoding="utf-8") as handle:
                json.dump({
                    "profile": {"name": "通用规则", "enabled_checks": []},
                    "rules": [],
                }, handle, ensure_ascii=False)

            result = check_document(document_path, rules_path)
            self.assertEqual(result.issues, [])

    def test_legacy_rule_list_keeps_previous_default_checks(self):
        with tempfile.TemporaryDirectory() as folder:
            document_path = self._document_path(folder)
            rules_path = os.path.join(folder, "legacy.json")
            with open(rules_path, "w", encoding="utf-8") as handle:
                json.dump([], handle)

            result = check_document(document_path, rules_path)
            self.assertGreater(len(result.issues), 0)

    def test_run_font_info_resolves_paragraph_style(self):
        doc = Document()
        style = doc.styles["Normal"]
        style.font.name = "SimSun"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.size = Pt(12)
        paragraph = doc.add_paragraph("继承样式")

        font_name, size_pt, _ = get_run_font_info(paragraph.runs[0])
        self.assertIn(font_name, ("宋体", "SimSun"))
        self.assertAlmostEqual(size_pt, 12.0, places=1)

    def test_combined_rule_keeps_font_and_line_spacing(self):
        with tempfile.TemporaryDirectory() as folder:
            rules_path = os.path.join(folder, "combined.json")
            with open(rules_path, "w", encoding="utf-8") as handle:
                json.dump([{
                    "section": "body_text",
                    "label": "正文",
                    "font": "宋体",
                    "size_pt": 12,
                    "line_spacing_pt": 22,
                }], handle, ensure_ascii=False)

            rules, _ = load_rules_from_json(rules_path)
            self.assertEqual(rules["body_text"]["font"], "宋体")
            self.assertEqual(rules["body_text"]["line_spacing_pt"], 22)

    def test_partial_custom_rules_do_not_require_default_thesis_fields(self):
        with tempfile.TemporaryDirectory() as folder:
            document_path = self._document_path(folder)
            rules_path = os.path.join(folder, "partial.json")
            with open(rules_path, "w", encoding="utf-8") as handle:
                json.dump({
                    "profile": {"enabled_checks": ["font"]},
                    "rules": [
                        {"section": "body_text", "label": "正文", "alignment": "左对齐"},
                        {"section": "cover_title", "label": "封面", "font": "宋体"},
                    ],
                }, handle, ensure_ascii=False)

            result = check_document(document_path, rules_path)
            self.assertIsNotNone(result.summary)


if __name__ == "__main__":
    unittest.main()
