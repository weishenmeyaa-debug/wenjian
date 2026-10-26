"""针对已修复缺陷的回归测试。

每个测试都对应一个曾经真实存在的错误行为，用于防止改动再次引入。
"""

import io
import json
import os
import tempfile
import unittest

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

import app as app_module
from checker import (
    check_document,
    check_figure_table_numbering,
    check_line_spacing,
    classify_paragraph,
)


def make_paragraph(doc, text, size=12, font="宋体", bold=False, align=None):
    """向文档追加一段，同时设置字号与东亚字体名。"""
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = rPr.makeelement(qn('w:rFonts'), {})
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), font)
    rFonts.set(qn('w:ascii'), font)
    if align is not None:
        paragraph.alignment = align
    return paragraph


def classify_all(doc, max_declaration_streak=True):
    """按 check_document 的方式对整篇文档做分类。"""
    classified = []
    prev = None
    streak = 0
    for para in doc.paragraphs:
        cls = classify_paragraph(para, prev, streak)
        classified.append(cls)
        if cls:
            prev = cls
        if cls in ("declaration_title", "declaration_body"):
            streak += 1
        else:
            streak = 0
    return classified


class DeclarationStickinessTests(unittest.TestCase):
    """声明区曾经会「粘住」整篇文档，导致摘要与关键词被吞掉。"""

    def _document(self):
        doc = Document()
        make_paragraph(doc, "诚信声明书", size=16, bold=True)
        make_paragraph(doc, "本人郑重声明：所呈交的材料由本人独立完成。", size=12)
        make_paragraph(doc, "签名：", size=12)
        make_paragraph(doc, "摘    要", size=15, font="黑体", bold=True)
        make_paragraph(doc, "本文研究了短视频内容的传播机制与效果。", size=12, font="楷体")
        make_paragraph(doc, "关键词：短视频；传播；用户行为", size=12, font="黑体")
        make_paragraph(doc, "1 绪论", size=15, bold=True)
        make_paragraph(doc, "这是正文段落的内容。", size=12)
        return doc

    def test_spaced_abstract_heading_is_classified(self):
        """「摘    要」这类带排版空格的标题必须能被识别。"""
        doc = self._document()
        classified = classify_all(doc)
        texts = [p.text.strip() for p in doc.paragraphs]

        self.assertIn("abstract_zh_title", classified,
                      f"分类结果异常：{list(zip(texts, classified))}")
        self.assertIn("abstract_zh_body", classified)
        self.assertIn("keywords_zh", classified)

    def test_declaration_stops_at_section_start(self):
        """声明区遇到新章节必须结束，不能无限延伸。"""
        doc = self._document()
        classified = classify_all(doc)
        texts = [p.text.strip() for p in doc.paragraphs]
        paired = list(zip(texts, classified))

        # 「签名：」仍属声明正文
        self.assertEqual(dict(paired)["签名："], "declaration_body")
        # 但「摘    要」之后不再有 declaration_body
        abstract_index = texts.index("摘    要")
        remaining = classified[abstract_index:]
        self.assertNotIn("declaration_body", remaining)

    def test_declaration_is_bounded_even_without_keywords(self):
        """即使没有任何章节关键词，声明区也不能吞掉整篇文档。"""
        doc = Document()
        make_paragraph(doc, "诚信声明书", size=16, bold=True)
        for i in range(30):
            make_paragraph(doc, f"正文段落 {i}，用于验证声明区不会无限延伸。", size=12)

        classified = classify_all(doc)
        declaration_count = sum(1 for c in classified if c == "declaration_body")
        self.assertLess(declaration_count, 30,
                        "声明区长度应当受上限约束")


class FigureTableNumberingTests(unittest.TestCase):
    """图/表序号连续性曾经只记录不比对，跳号不会报错。"""

    def _run(self, texts):
        doc = Document()
        for t in texts:
            make_paragraph(doc, t, size=10.5)
        paragraphs = doc.paragraphs
        from checker import check_figure_table_numbering
        issues = []
        check_figure_table_numbering(paragraphs, ["figure_caption"] * len(paragraphs), issues)
        return issues

    def test_figure_gap_is_reported(self):
        issues = self._run(["图3-1 增长趋势", "图3-5 另一张图"])
        self.assertTrue(
            any(i.severity.value == "错误" and "不连续" in i.description for i in issues),
            f"图编号跳号应当被报错，实际：{[i.description for i in issues]}",
        )

    def test_sequential_figures_are_clean(self):
        issues = self._run(["图3-1 第一张", "图3-2 第二张", "图3-3 第三张"])
        self.assertEqual(
            [i.description for i in issues if "不连续" in i.description], [],
            "连续编号不应报错",
        )

    def test_subfigures_share_one_number(self):
        """子图 (a)(b) 共用主编号，不应被当成跳号。"""
        issues = self._run([
            "图2-1 短视频封面（a）",
            "图2-1 短视频封面（b）",
            "图2-2 用户画像",
        ])
        self.assertEqual(
            [i.description for i in issues if "不连续" in i.description], [],
            f"子图不应触发跳号，实际：{[i.description for i in issues]}",
        )


class LineSpacingRuleTests(unittest.TestCase):
    """自定义模板下，参考文献行距分支曾经是永不执行的死代码。"""

    def test_reference_line_spacing_is_checked(self):
        doc = Document()
        paragraph = make_paragraph(doc, "[1] 王五. 短视频研究[J]. 新闻界, 2024(3): 12-15.", size=10.5)
        paragraph.paragraph_format.line_spacing = Pt(30)

        rules = {"reference_entry": {"line_spacing_pt": 18, "line_spacing_type": "固定值"}}
        issues = []
        check_line_spacing("reference_entry", doc.paragraphs[0], issues, 0, rules)

        self.assertTrue(
            any(i.category == "行间距" for i in issues),
            "参考文献行距偏离规则时应当报错",
        )


class FootnoteRobustnessTests(unittest.TestCase):
    """脚注检查曾经假设 id 从 1 起且必然是数字，导致误报或崩溃。"""

    def test_non_numeric_footnote_id_does_not_crash(self):
        from lxml import etree

        from checker import check_footnotes

        doc = Document()
        make_paragraph(doc, "正文内容。", size=12)

        # 手工插入一个非数字编号的脚注引用，模拟异常文档
        body = doc.element.body
        paragraph = doc.paragraphs[0]._element
        ref = etree.SubElement(paragraph, qn('w:r'))
        ref_child = etree.SubElement(ref, qn('w:footnoteReference'))
        ref_child.set(qn('w:id'), "abc")

        try:
            issues = []
            check_footnotes(doc, doc.paragraphs, ["body_text"], issues)
        except Exception as exc:  # noqa: BLE001 - 回归测试就是要捕获任何崩溃
            self.fail(f"非数字脚注编号不应导致崩溃：{exc}")


class PathSafetyTests(unittest.TestCase):
    """模板标识解析曾经会把无法识别的标识兜底成内置模板。"""

    def test_unknown_profile_does_not_fall_back(self):
        self.assertIsNone(app_module.resolve_profile(".."))
        self.assertIsNone(app_module.resolve_profile("不存在的文件.json"))
        self.assertIsNone(app_module.resolve_profile(""))

    def test_known_builtin_profile_resolves(self):
        profile = app_module.resolve_profile(app_module.BUILTIN_PROFILE_ID)
        self.assertIsNotNone(profile)
        self.assertTrue(profile["is_builtin"])

    def test_editor_load_rejects_traversal(self):
        client = app_module.app.test_client()
        for candidate in ("..", "%2E%2E", "%2E%2E%2Fformat_rules.json"):
            response = client.get("/editor/load/" + candidate)
            self.assertEqual(
                response.status_code, 404,
                f"路径穿越标识必须返回 404：{candidate}",
            )

    def test_editor_delete_rejects_builtin(self):
        client = app_module.app.test_client()
        response = client.post("/editor/delete/__builtin__")
        self.assertEqual(response.status_code, 400)
        self.assertTrue(os.path.exists(os.path.join(app_module.BASE_DIR, "format_rules.json")))


class UploadLimitTests(unittest.TestCase):
    """超过体积限制时，曾经直接返回 Flask 默认错误页。"""

    def test_oversized_upload_redirects_with_flash(self):
        app_module.app.config.update(TESTING=True)
        client = app_module.app.test_client()
        limit = app_module.app.config["MAX_CONTENT_LENGTH"]

        response = client.post(
            "/check",
            data={
                "files": (io.BytesIO(b"x" * (limit + 1024)), "big.docx"),
                "profile_id": app_module.BUILTIN_PROFILE_ID,
            },
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.headers.get("Location"), "/")


class ProductNamingTests(unittest.TestCase):
    """公开页面不得出现特定材料名称，保持工具的中性定位。"""

    def setUp(self):
        app_module.app.config.update(TESTING=True)
        self.client = app_module.app.test_client()

    def test_public_pages_use_neutral_name(self):
        for path in ("/", "/editor", "/examples"):
            text = self.client.get(path).get_data(as_text=True)
            self.assertIn("格式检查器", text, f"{path} 应显示产品名")
            for word in ("毕业设计", "毕业论文", "本科"):
                self.assertNotIn(word, text,
                                 f"{path} 不应出现「{word}」")


if __name__ == "__main__":
    unittest.main()
